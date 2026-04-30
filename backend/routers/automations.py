from io import BytesIO
from uuid import UUID
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from database import get_db
from models import Automation, User
from schemas import AutomationOut
from auth_deps import get_current_user

router = APIRouter(tags=["automations"])


class AutomationCreate(BaseModel):
    category: str
    sub_category: str = ""
    script_name: str
    description: Optional[str] = None

_HEADERS = ["Category", "Sub-Category", "ScriptName", "Description"]


@router.get("/automations", response_model=list[AutomationOut])
async def list_automations(
    category: Optional[str] = Query(None),
    sub_category: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    stmt = select(Automation).order_by(Automation.created_at)
    if category:
        stmt = stmt.where(Automation.category == category)
    if sub_category:
        stmt = stmt.where(Automation.sub_category == sub_category)
    if search:
        term = f"%{search}%"
        stmt = stmt.where(
            Automation.script_name.ilike(term) | Automation.description.ilike(term)
        )
    result = await db.execute(stmt)
    return result.scalars().all()


@router.post("/automations", response_model=AutomationOut, status_code=status.HTTP_201_CREATED)
async def create_automation(
    body: AutomationCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    if not body.category.strip():
        raise HTTPException(status_code=422, detail="category is required")
    if not body.script_name.strip():
        raise HTTPException(status_code=422, detail="script_name is required")
    obj = Automation(
        category=body.category.strip(),
        sub_category=body.sub_category.strip(),
        script_name=body.script_name.strip(),
        description=body.description.strip() if body.description else None,
    )
    db.add(obj)
    await db.commit()
    await db.refresh(obj)
    return obj


@router.get("/automations/template")
async def download_template(current_user: User = Depends(get_current_user)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Automations"
    ws.append(_HEADERS)
    ws.append(["COMPUTE", "UserManagement", "COMPUTE_POWERSHELL_ADD-USER-IN-DL", "Example description"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="automations_template.xlsx"'},
    )


@router.post("/automations/upload", response_model=list[AutomationOut], status_code=status.HTTP_201_CREATED)
async def upload_automations(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid Excel file (.xlsx required)")

    ws = wb.active
    created: list[Automation] = []

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not row or all(v is None for v in row):
            continue
        if len(row) < 3:
            continue
        category, sub_category, script_name = row[0], row[1], row[2]
        description = row[3] if len(row) > 3 else None

        if not category or not script_name:
            continue

        obj = Automation(
            category=str(category).strip(),
            sub_category=str(sub_category).strip() if sub_category else "",
            script_name=str(script_name).strip(),
            description=str(description).strip() if description else None,
        )
        db.add(obj)
        created.append(obj)

    if not created:
        raise HTTPException(status_code=422, detail="No valid rows found in the uploaded file")

    await db.commit()
    for obj in created:
        await db.refresh(obj)
    return created


@router.delete("/automations/{automation_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_automation(
    automation_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(Automation).where(Automation.id == automation_id))
    obj = result.scalar_one_or_none()
    if not obj:
        raise HTTPException(status_code=404, detail="Automation not found")
    await db.delete(obj)
    await db.commit()
